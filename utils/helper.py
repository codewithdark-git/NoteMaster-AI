from fpdf import FPDF
from docx import Document
import io
import tempfile
import streamlit as st
import pyperclip


def save_btn(bot_response):
    col1, col2, col3, col4, col5 = st.columns(5)
    with col1:
        if st.download_button('Save as TXT', save_as_txt(bot_response), file_name='notes.txt'):
            st.toast('Notes Save as TXT', icon='🎉')
    with col2:
        if st.download_button('Save as PDF', save_as_pdf(bot_response), file_name='notes.pdf'):
            st.toast('Notes Save as PDF', icon='🎉')
    with col3:
        if st.download_button('Save as DOCX', save_as_doc(bot_response), file_name='notes.docx'):
            st.toast('Notes Save as DOCX', icon='🎉')
    with col4:
        if st.button('Copy'):
            pyperclip.copy(bot_response)
            st.toast('Notes Copy to Clipboard', icon='🎉')
    st.markdown('---')


def save_as_txt(content):
    buffer = io.BytesIO()
    buffer.write(content.encode())
    buffer.seek(0)
    return buffer


def save_as_pdf(content):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    lines = content.split('\n')
    for line in lines:
        pdf.cell(200, 10, txt=line, ln=True)

    # Save to a temporary file
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_file:
        pdf.output(tmp_file.name)
        tmp_file.seek(0)
        pdf_content = tmp_file.read()

    return pdf_content


def save_as_doc(content):
    doc = Document()
    doc.add_paragraph(content)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer



