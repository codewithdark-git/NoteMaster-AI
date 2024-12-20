import logging
from fpdf import FPDF
from docx import Document
from PIL import Image
import io
import tempfile
import streamlit as st
import numpy as np
from backend.utils.ocr import image_to_text
import pyperclip

# Configure logging
logging.basicConfig(level=logging.DEBUG)
logger = logging.getLogger(__name__)

def save_btn(bot_response):
    col1, col2, col3, col4 = st.columns(4)
    with col1:
        if st.download_button('Save as TXT', save_as_txt(bot_response), file_name='notes.txt'):
            st.toast('Notes Save as TXT', icon='🎉')
    with col2:
        if st.download_button('Save as PDF', save_as_pdf(bot_response), file_name='notes.pdf'):
            st.toast('Notes Save as PDF', icon='🎉')
    with col3:
        if st.download_button('Save as DOCX', save_as_doc(bot_response), file_name='notes.docx'):
            st.toast('Notes Save as DOCX', icon='🎉')
    with col4:
        if st.button('Copy'):
            pyperclip.copy(bot_response)
            st.toast('Notes Copy to Clipboard', icon='🎉')
    st.markdown('---')

def process_images(uploaded_files):
    images = []
    text_list = []
    try:
        
        for uploaded_file in uploaded_files:
            try:  
                # Open image from BytesIO
                image = Image.open(uploaded_file)
                # Convert to RGB if necessary
                if image.mode == 'RGBA':
                    image = image.convert('RGB')
                    
                # Convert to numpy array for OCR
                image_array = np.array(image)
                 
                # Extract text using OCR
                extracted_text = image_to_text(image_array)
                
                # Only append non-empty text
                if extracted_text and isinstance(extracted_text, str) and extracted_text.strip():
                    text_list.append(extracted_text)
                    images.append(image)
                    
            except Exception as img_error:
                    continue
        
        if not text_list:
            
            raise Exception("No text could be extracted from any of the images")     
        return text_list, images
        
    except Exception as e:

        raise Exception(f"Error processing images: {str(e)}")
    finally:
        # Clean up BytesIO objects
        for file in uploaded_files:
            try:
                file.close()
            except Exception as close_error:
                pass

def save_as_txt(content):
    buffer = io.BytesIO()
    buffer.write(content.encode())
    buffer.seek(0)
    return buffer


def save_as_pdf(content):
    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)

        lines = content.split('\n')
        for line in lines:
            pdf.cell(200, 10, txt=line, ln=True)

        # Save to a temporary file
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_file:
            pdf.output(tmp_file.name)
            tmp_file.seek(0)
            pdf_content = tmp_file.read()

        return pdf_content

    except Exception as e:
        return f"Error Occur {e}"

def save_as_doc(content):
    doc = Document()
    doc.add_paragraph(content)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
